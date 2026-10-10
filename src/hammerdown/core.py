from __future__ import annotations

import logging
import os
from pathlib import Path
import time
from hammerdown.parsers import CONVERTERS
from hammerdown.parsers.charts import render_chart_svg
from hammerdown.parsers.math import mathml_to_latex, omml_to_latex
from hammerdown.parsers.office import (
    convert_doc,
    convert_docx,
    convert_odt,
    convert_ppt,
    convert_pptx,
    convert_xls,
    convert_xlsx,
)
from hammerdown.parsers.pdf import convert_pdf_or_ebook
from hammerdown.parsers.tables import render_table_regions
from hammerdown.parsers.text import convert_csv, convert_txt_or_md
from hammerdown.utils import format_duration, normalize_path

logger = logging.getLogger("hammerdown")

SUPPORTED_EXTENSIONS: tuple[str, ...] = tuple(CONVERTERS)


def to_markdown(
    file_path: str | os.PathLike[str],
    out_dir: str | os.PathLike[str] | None = None,
) -> tuple[str | None, int]:
    normalized_path = normalize_path(file_path)
    source = Path(normalized_path)
    if not source.is_file():
        logger.error("File not found: %s", normalized_path)
        return None, 0

    extension = source.suffix.lower()
    converter = CONVERTERS.get(extension)
    if converter is None:
        logger.error("Unsupported format: %s", extension or "(no extension)")
        return None, 0

    target_dir = str(out_dir) if out_dir is not None else str(source.parent)
    return converter(str(source), target_dir)


def convert_file(
    file_path: str | os.PathLike[str],
    force: bool = False,
    in_place: bool = True,
) -> bool:
    normalized_path = normalize_path(file_path)
    source = Path(normalized_path)
    if not source.is_file():
        logger.error("File not found: %s", normalized_path)
        return False

    stem = source.stem
    extension = source.suffix.lower()
    ext_clean = extension.lstrip(".")

    out_dir = source.parent
    out_md_name = f"{stem}_{ext_clean}.md" if ext_clean else f"{stem}.md"
    out_md = out_dir / out_md_name

    if out_md.exists() and not force:
        logger.error("Destination file already exists: %s. Use --force to overwrite.", out_md)
        return False

    out_dir.mkdir(parents=True, exist_ok=True)

    converter = CONVERTERS.get(extension)
    if converter is None:
        logger.error("Unsupported format: %s", extension or "(no extension)")
        return False

    est_hint = ""
    if extension in {".pdf", ".epub", ".mobi", ".fb2", ".xps"}:
        try:
            import pymupdf  # type: ignore
            with pymupdf.open(str(source)) as doc:
                page_count = len(doc)
            est_seconds = max(1.0, page_count * 0.9)
            est_hint = f" ({page_count} pages, est. ~{format_duration(est_seconds)})"
        except Exception:
            pass
    elif extension in {".pptx"}:
        try:
            from pptx import Presentation  # type: ignore
            prs = Presentation(str(source))
            slide_count = len(prs.slides)
            est_seconds = max(1.0, slide_count * 0.1)
            est_hint = f" ({slide_count} slides, est. ~{format_duration(est_seconds)})"
        except Exception:
            pass

    logger.info("Processing %s%s...", source.name, est_hint)
    started_at = time.monotonic()
    try:
        md_text, saved_images = converter(str(source), str(out_dir))
        if md_text is None:
            return False

        # Trim trailing whitespace on each line and ensure a single final newline
        normalized_text = md_text.replace("\r\n", "\n").replace("\r", "\n")
        trimmed_lines = [line.rstrip() for line in normalized_text.split("\n")]
        cleaned_md = "\n".join(trimmed_lines).rstrip() + "\n"
        out_md.write_text(cleaned_md, encoding="utf-8", newline="\n")
        duration = time.monotonic() - started_at
        logger.info(
            "Completed %s in %s (%d images extracted)",
            source.name,
            format_duration(duration),
            saved_images,
        )
        return True
    except Exception:
        logger.exception("Failed to convert %s", source.name)
        return False
