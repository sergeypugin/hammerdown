from __future__ import annotations

import logging
from typing import Callable

from hammerdown.parsers.markdown_render import render_markdown_table

logger = logging.getLogger("hammerdown")
ProgressCallback = Callable[[int, int, str], None]


def convert_pptx(
    pptx_path: str,
    out_dir: str,
    progress_callback: ProgressCallback | None = None,
) -> tuple[str | None, int]:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        logger.error("PPTX support requires python-pptx")
        return None, 0

    presentation = Presentation(pptx_path)
    lines: list[str] = []
    slides_list = list(presentation.slides)
    total_slides = len(slides_list)
    if progress_callback:
        progress_callback(0, total_slides, "slides")

    for slide_idx, slide in enumerate(slides_list, start=1):
        lines.append(f"## Slide {slide_idx}\n")
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                table = getattr(shape, "table", None)
                if table is None:
                    continue
                rows = list(table.rows)
                if rows:
                    rendered_rows = [[cell.text.strip() for cell in row.cells] for row in rows]
                    lines.extend(render_markdown_table(rendered_rows))
            elif getattr(shape, "has_text_frame", False):
                text_frame = getattr(shape, "text_frame", None)
                if text_frame is None:
                    continue
                for paragraph in text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        lines.append(text)
        lines.append("")
        if progress_callback:
            progress_callback(slide_idx, total_slides, "slides")

    return "\n".join(lines), 0
