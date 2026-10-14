from __future__ import annotations

import logging
from typing import Callable

from hammerdown.parsers.tables import render_table_regions

logger = logging.getLogger("hammerdown")
ProgressCallback = Callable[[int, int, str], None]


def convert_xlsx(
    xlsx_path: str,
    out_dir: str,
    progress_callback: ProgressCallback | None = None,
) -> tuple[str | None, int]:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        logger.error("XLSX support requires openpyxl")
        return None, 0

    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    formula_workbook = openpyxl.load_workbook(xlsx_path, data_only=False, read_only=True)
    blocks: list[str] = []
    total_sheets = len(workbook.sheetnames)
    if progress_callback:
        progress_callback(0, total_sheets, "sheets")

    for sheet_idx, sheet in enumerate(workbook.sheetnames, start=1):
        blocks.append(f"# Sheet: {sheet}")
        data_sheet = workbook[sheet]
        formula_sheet = formula_workbook[sheet]
        rows = []
        for data_row, formula_row in zip(data_sheet.iter_rows(), formula_sheet.iter_rows()):
            row = []
            for data_cell, formula_cell in zip(data_row, formula_row):
                value = data_cell.value
                if value is None and formula_cell.data_type == "f":
                    value = formula_cell.value
                row.append(value)
            rows.append(row)
        blocks.extend(render_table_regions(rows))
        if progress_callback:
            progress_callback(sheet_idx, total_sheets, "sheets")
    workbook.close()
    formula_workbook.close()
    return "\n\n".join(blocks), 0
