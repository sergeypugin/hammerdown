from __future__ import annotations

import logging
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile

from hammerdown.parsers.charts import render_chart_svg
from hammerdown.parsers.math import MATH_NS as _MATH_NS
from hammerdown.parsers.math import omml_to_latex as _omml_to_latex
from hammerdown.parsers.math import xml_name as _xml_name
from hammerdown.parsers.odt import convert_odt
from hammerdown.parsers.tables import render_table_regions
from hammerdown.utils import get_images_dir_name

logger = logging.getLogger("hammerdown")


def _convert_via_soffice(file_path: str, out_ext: str) -> str | None:
    soffice_bin = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice_bin and os.name == "nt":
        common_paths = [
            Path(os.environ.get("PROGRAMFILES", "C:\\Program Files")) / "LibreOffice" / "program" / "soffice.exe",
            Path(os.environ.get("PROGRAMFILES(X86)", "C:\\Program Files (x86)")) / "LibreOffice" / "program" / "soffice.exe",
        ]
        for p in common_paths:
            if p.is_file():
                soffice_bin = str(p)
                break

    if not soffice_bin:
        return None

    temp_dir = tempfile.mkdtemp()
    cmd = [
        soffice_bin,
        "--headless",
        "--convert-to",
        out_ext,
        "--outdir",
        temp_dir,
        file_path,
    ]
    try:
        subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        converted_file = Path(temp_dir) / f"{Path(file_path).stem}.{out_ext}"
        if converted_file.is_file():
            return str(converted_file)
    except Exception:
        pass
    return None


def _render_word_node(node, relationships, chart_links=None) -> str | tuple[str, str]:
    tag = _xml_name(node)
    if tag == "chart":
        relation_id = node.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        return chart_links.get(relation_id, "") if chart_links else ""
    if tag == "t":
        return node.text or ""
    if tag == "tab":
        return "\t"
    if tag in {"br", "cr"}:
        return "\n"
    if tag == "oMath":
        return ("inline", _omml_to_latex(node).strip())
    if tag == "oMathPara":
        formula = " ".join(_omml_to_latex(math).strip() for math in node.findall(f".//{_MATH_NS}oMath"))
        return ("display", formula.strip())
    if tag == "hyperlink":
        rendered = _render_word_children(node, relationships, chart_links)
        relation_id = node.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        url = relationships[relation_id].target_ref if relation_id in relationships else ""
        if url and rendered and isinstance(rendered, str):
            return f"[{rendered}]({url})"
        return rendered
    if tag.endswith("Pr") or tag in {"pPr", "sectPr"}:
        return ""
    return _render_word_children(node, relationships, chart_links)


def _render_word_children(node, relationships, chart_links=None) -> str | tuple[str, str]:
    parts = [_render_word_node(child, relationships, chart_links) for child in node]
    formulas = [part for part in parts if isinstance(part, tuple)]
    text = "".join(part for part in parts if isinstance(part, str))
    if not formulas:
        return text
    has_text = bool(text.strip())
    rendered = []
    for part in parts:
        if isinstance(part, tuple):
            mode, latex = part
            delimiter = "$$" if mode == "display" or not has_text else "$"
            prefix = " " if rendered and isinstance(rendered[-1], str) and rendered[-1] and not rendered[-1].endswith((" ", "\n", "\t", "(")) else ""
            rendered.append(f"{prefix}{delimiter}{latex}{delimiter}")
        else:
            rendered.append(part)
    return "".join(rendered)


def _render_word_paragraph(element, relationships, chart_links=None) -> str:
    rendered = _render_word_children(element, relationships, chart_links)
    return rendered.strip() if isinstance(rendered, str) else ""


def convert_docx(docx_path: str, out_dir: str, images_dir_name: str | None = None) -> tuple[str | None, int]:
    try:
        import docx  # type: ignore
        from docx.table import Table
    except ImportError:
        logger.error("DOCX support requires python-docx")
        return None, 0

    doc = docx.Document(docx_path)
    chart_links: dict[str, str] = {}
    saved_images = 0
    img_dir_name = images_dir_name or get_images_dir_name(docx_path)
    image_dir = Path(out_dir) / img_dir_name
    for relation_id, relation in doc.part.rels.items():
        if not relation.reltype.endswith("/chart"):
            continue
        rendered_chart = render_chart_svg(relation.target_part.blob)
        if rendered_chart is None:
            continue
        title, svg = rendered_chart
        image_dir.mkdir(parents=True, exist_ok=True)
        image_name = f"chart_{saved_images:03d}.svg"
        (image_dir / image_name).write_text(svg, encoding="utf-8")
        chart_links[relation_id] = f"![{title}]({img_dir_name}/{image_name})"
        saved_images += 1

    blocks: list[str] = []
    for element in doc.element.body:
        tag = _xml_name(element)
        if tag == "p":
            text = _render_word_paragraph(element, doc.part.rels, chart_links)
            if text:
                blocks.append(text)
        elif tag == "tbl":
            table = Table(element, doc)
            rows = []
            for row_index, row in enumerate(table.rows):
                cells = []
                for cell in row.cells:
                    cell_paragraphs = [
                        _render_word_paragraph(paragraph._p, doc.part.rels, chart_links)
                        for paragraph in cell.paragraphs
                    ]
                    cells.append(" ".join(text for text in cell_paragraphs if text).replace("|", "\\|"))
                rows.append(f"| {' | '.join(cells)} |")
                if row_index == 0:
                    rows.append(f"| {' | '.join(['---'] * len(cells))} |")
            if rows:
                blocks.append("\n".join(rows))

    return "\n\n".join(blocks), saved_images


def convert_xlsx(xlsx_path: str, out_dir: str) -> tuple[str | None, int]:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        logger.error("XLSX support requires openpyxl")
        return None, 0

    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    formula_workbook = openpyxl.load_workbook(xlsx_path, data_only=False, read_only=True)
    blocks: list[str] = []
    for sheet in workbook.sheetnames:
        blocks.append(f"# Sheet: {sheet}")
        data_sheet = workbook[sheet]
        formula_sheet = formula_workbook[sheet]
        rows = []
        for data_row, formula_row in zip(data_sheet.iter_rows(), formula_sheet.iter_rows()):
            row = []
            for data_cell, formula_cell in zip(data_row, formula_row):
                value = data_cell.value
                if value is None and formula_cell.data_type == "f":
                    value = formula_cell.value
                row.append(value)
            rows.append(row)
        blocks.extend(render_table_regions(rows))
    workbook.close()
    formula_workbook.close()
    return "\n\n".join(blocks), 0


def convert_pptx(pptx_path: str, out_dir: str) -> tuple[str | None, int]:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        logger.error("PPTX support requires python-pptx")
        return None, 0

    presentation = Presentation(pptx_path)
    lines: list[str] = []
    for slide_idx, slide in enumerate(list(presentation.slides), start=1):
        lines.append(f"## Slide {slide_idx}\n")
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                table = getattr(shape, "table", None)
                if table is None:
                    continue
                rows = list(table.rows)
                if rows:
                    header = [cell.text.strip() for cell in rows[0].cells]
                    lines.append(f"| {' | '.join(header)} |")
                    lines.append(f"| {' | '.join(['---'] * len(header))} |")
                    for row in rows[1:]:
                        lines.append(f"| {' | '.join(cell.text.strip() for cell in row.cells)} |")
            elif getattr(shape, "has_text_frame", False):
                text_frame = getattr(shape, "text_frame", None)
                if text_frame is None:
                    continue
                for paragraph in text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        lines.append(text)
        lines.append("")

    return "\n".join(lines), 0


def convert_doc(doc_path: str, out_dir: str) -> tuple[str | None, int]:
    img_dir_name = get_images_dir_name(doc_path)
    if zipfile.is_zipfile(doc_path):
        return convert_docx(doc_path, out_dir, images_dir_name=img_dir_name)

    converted = _convert_via_soffice(doc_path, "docx")
    if converted:
        return convert_docx(converted, out_dir, images_dir_name=img_dir_name)

    if os.name == "nt":
        try:
            import win32com.client  # type: ignore
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            temp_file = Path(tempfile.mkdtemp()) / f"{Path(doc_path).stem}.docx"
            doc = word.Documents.Open(str(Path(doc_path).resolve()))
            doc.SaveAs2(str(temp_file.resolve()), FileFormat=16)
            doc.Close()
            word.Quit()
            if temp_file.is_file():
                return convert_docx(str(temp_file), out_dir, images_dir_name=img_dir_name)
        except Exception:
            pass

    logger.error("Legacy .doc format requires LibreOffice or Microsoft Word to be installed")
    return None, 0


def convert_xls(xls_path: str, out_dir: str) -> tuple[str | None, int]:
    if zipfile.is_zipfile(xls_path):
        return convert_xlsx(xls_path, out_dir)

    converted = _convert_via_soffice(xls_path, "xlsx")
    if converted:
        return convert_xlsx(converted, out_dir)

    if os.name == "nt":
        try:
            import win32com.client  # type: ignore
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = False
            temp_file = Path(tempfile.mkdtemp()) / f"{Path(xls_path).stem}.xlsx"
            wb = excel.Workbooks.Open(str(Path(xls_path).resolve()))
            wb.SaveAs(str(temp_file.resolve()), FileFormat=51)
            wb.Close()
            excel.Quit()
            if temp_file.is_file():
                return convert_xlsx(str(temp_file), out_dir)
        except Exception:
            pass

    logger.error("Legacy .xls format requires LibreOffice or Microsoft Excel to be installed")
    return None, 0


def convert_ppt(ppt_path: str, out_dir: str) -> tuple[str | None, int]:
    if zipfile.is_zipfile(ppt_path):
        return convert_pptx(ppt_path, out_dir)

    converted = _convert_via_soffice(ppt_path, "pptx")
    if converted:
        return convert_pptx(converted, out_dir)

    if os.name == "nt":
        try:
            import win32com.client  # type: ignore
            powerpoint = win32com.client.Dispatch("PowerPoint.Application")
            temp_file = Path(tempfile.mkdtemp()) / f"{Path(ppt_path).stem}.pptx"
            ppt = powerpoint.Presentations.Open(str(Path(ppt_path).resolve()), WithWindow=False)
            ppt.SaveAs(str(temp_file.resolve()), FileFormat=24)
            ppt.Close()
            powerpoint.Quit()
            if temp_file.is_file():
                return convert_pptx(str(temp_file), out_dir)
        except Exception:
            pass

    logger.error("Legacy .ppt format requires LibreOffice or Microsoft PowerPoint to be installed")
    return None, 0


CONVERTERS = {
    ".docx": convert_docx,
    ".doc": convert_doc,
    ".odt": convert_odt,
    ".xlsx": convert_xlsx,
    ".xls": convert_xls,
    ".pptx": convert_pptx,
    ".ppt": convert_ppt,
}
