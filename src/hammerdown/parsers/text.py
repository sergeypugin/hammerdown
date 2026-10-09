from __future__ import annotations

import base64
import csv
import io
import re
from pathlib import Path
import zipfile


def convert_txt_or_md(file_path: str, out_dir: str) -> tuple[str | None, int]:
    raw = Path(file_path).read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1251", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")

    # Extract base64 images from Markdown/text
    # Pattern: data:image/(png|jpeg|jpg|webp|gif);base64,[data]
    img_dir = Path(out_dir) / "images"
    saved_imgs = 0

    def replacer(match: re.Match[str]) -> str:
        nonlocal saved_imgs
        ext = match.group(1)
        b64_data = match.group(2)
        try:
            img_bytes = base64.b64decode(b64_data)
            if not img_dir.exists():
                img_dir.mkdir(parents=True, exist_ok=True)

            img_name = f"embedded_img_{saved_imgs}.{ext}"
            (img_dir / img_name).write_bytes(img_bytes)
            saved_imgs += 1
            return f"images/{img_name}"
        except Exception:
            return match.group(0)

    # Replace data URIs in Markdown image syntax or raw links
    # Matches: data:image/png;base64,iVBOR...
    pattern = r"data:image/([a-zA-Z]+);base64,([a-zA-Z0-9+/=]+)"
    processed_text = re.sub(pattern, replacer, text)

    return processed_text, saved_imgs


def convert_csv(csv_path: str, out_dir: str) -> tuple[str | None, int]:
    if zipfile.is_zipfile(csv_path):
        try:
            import openpyxl  # type: ignore
            data = Path(csv_path).read_bytes()
            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            lines: list[str] = []
            for sheet in workbook.sheetnames:
                lines.append(f"# Sheet: {sheet}\n")
                rows = workbook[sheet].iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    continue
                header_cells = [str(cell) if cell is not None else "" for cell in header]
                lines.append(f"| {' | '.join(header_cells)} |")
                lines.append(f"| {' | '.join(['---'] * len(header_cells))} |")
                for row in rows:
                    if any(row):
                        row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                        lines.append(f"| {row_text} |")
                lines.append("\n")
            workbook.close()
            return "\n".join(lines), 0
        except Exception:
            pass

    raw = Path(csv_path).read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1251", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        text = raw.decode("utf-8", errors="replace")

    sample = text[:2048]
    delimiter = ","
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except Exception:
        if ";" in sample and "," not in sample:
            delimiter = ";"
        elif "\t" in sample and "," not in sample:
            delimiter = "\t"

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return "", 0

    lines = []
    header = [cell.strip() for cell in rows[0]]
    lines.append(f"| {' | '.join(header)} |")
    lines.append(f"| {' | '.join(['---'] * len(header))} |")
    for row in rows[1:]:
        if any(cell.strip() for cell in row):
            cells = [cell.strip() for cell in row]
            lines.append(f"| {' | '.join(cells)} |")

    return "\n".join(lines), 0
