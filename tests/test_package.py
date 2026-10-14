from pathlib import Path

import hammerdown
from hammerdown.cli import main
from hammerdown.core import convert_file, to_markdown
from hammerdown.utils import normalize_path


def test_package_exports():
    assert hasattr(hammerdown, "to_markdown")
    assert hasattr(hammerdown, "convert_file")
    assert hasattr(hammerdown, "convert_docx")
    assert hasattr(hammerdown, "convert_doc")
    assert hasattr(hammerdown, "convert_odt")
    assert hasattr(hammerdown, "convert_xlsx")
    assert hasattr(hammerdown, "convert_xls")
    assert hasattr(hammerdown, "convert_pptx")
    assert hasattr(hammerdown, "convert_ppt")
    assert hasattr(hammerdown, "convert_pdf_or_ebook")
    assert hasattr(hammerdown, "convert_csv")
    assert hasattr(hammerdown, "convert_txt_or_md")
    assert hasattr(hammerdown, "render_table_regions")
    assert hasattr(hammerdown, "render_chart_svg")
    assert hasattr(hammerdown, "omml_to_latex")
    assert hasattr(hammerdown, "mathml_to_latex")
    assert hasattr(hammerdown, "SUPPORTED_EXTENSIONS")
    assert hasattr(hammerdown, "TEXT_EXTENSIONS")
    assert hasattr(hammerdown, "APP_NAME")
    assert hasattr(hammerdown, "APP_DESCRIPTION")
    assert hasattr(hammerdown, "DEFAULT_IN_PLACE")
    assert hasattr(hammerdown, "DEFAULT_OVERWRITE")
    assert hasattr(hammerdown, "DEFAULT_SKIP_EXISTING")
    assert hasattr(hammerdown, "DEFAULT_QUIET")
    assert hasattr(hammerdown, "__version__")


def test_to_markdown_plain_text(tmp_path: Path):
    source = tmp_path / "hello.txt"
    source.write_text("Hello world!", encoding="utf-8")

    text, saved_images = to_markdown(source)
    assert text == "Hello world!"
    assert saved_images == 0


def test_convert_file_package(tmp_path: Path):
    source = tmp_path / "doc.txt"
    source.write_text("package test", encoding="utf-8")

    assert convert_file(source)
    out_file = tmp_path / "hm_doc_txt.md"
    assert out_file.is_file()
    assert out_file.read_text(encoding="utf-8") == "package test\n"


def test_batch_ui_error_panel_keeps_fixed_height(tmp_path: Path):
    from io import StringIO

    from rich.console import Console

    from hammerdown.utils import BATCH_UI_WIDTH, BatchUI

    ui = BatchUI([str(tmp_path / "one.txt"), str(tmp_path / "two.txt")])
    console = Console(file=StringIO(), width=120, force_terminal=False)
    before_renderable = ui._renderable()
    before = console.render_lines(before_renderable)
    ui.errors.append(("one.txt", "missing file " * 80))
    after_renderable = ui._renderable()
    after = console.render_lines(after_renderable)

    assert console.measure(before_renderable).maximum == BATCH_UI_WIDTH
    assert console.measure(after_renderable).maximum == BATCH_UI_WIDTH
    assert len(before) == len(after)


def test_batch_ui_completion_logs_token_count(tmp_path: Path, caplog):
    import logging
    from hammerdown.utils import BatchUI

    ui = BatchUI([str(tmp_path / "document.pdf")])
    ui.is_tty = False

    with caplog.at_level(logging.INFO, logger="hammerdown"):
        ui.complete_file(0, True, "5s", images=2, tokens=123)

    assert "Completed document.pdf (5s, 2 images, ~123 tokens)" in caplog.text


def test_progress_callback(tmp_path: Path):
    from hammerdown.utils import format_progress_bar

    bar = format_progress_bar(5, 10, unit="pages", bar_length=10)
    assert "[█████░░░░░] 50% (5/10 pages)" == bar

    source = tmp_path / "report.docx"
    import docx
    doc = docx.Document()
    doc.add_paragraph("Paragraph 1")
    doc.add_paragraph("Paragraph 2")
    doc.save(str(source))

    calls: list[tuple[int, int, str]] = []

    def on_progress(current: int, total: int, unit: str):
        calls.append((current, total, unit))

    text, images = to_markdown(source, progress_callback=on_progress)
    assert text is not None
    assert len(calls) > 0
    assert calls[-1][0] == calls[-1][1]  # final call reaches 100%
