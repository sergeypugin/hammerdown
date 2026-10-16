from pathlib import Path

import pytest

import hammerdown.cli
import hammerdown.core
import hammerdown.utils
import os

def test_keyboard_interrupt_is_reported_as_cancellation(monkeypatch, capsys):
    def interrupt(_argv):
        raise KeyboardInterrupt

    monkeypatch.setattr(hammerdown.cli, "_main", interrupt)

    assert hammerdown.cli.main([]) == 130
    assert "Cancelled." in capsys.readouterr().err


def test_normalize_wsl_windows_path():
    normalized = hammerdown.utils.normalize_path(r"C:\Users\someone\document.pdf")
    if os.name == "nt":
        assert normalized.endswith(r"C:\Users\someone\document.pdf")
    else:
        assert normalized.endswith("/mnt/c/Users/someone/document.pdf")


def test_convert_plain_text_creates_markdown_output(tmp_path: Path):
    source = tmp_path / "notes.txt"
    source.write_text("line one\nline two\n", encoding="utf-8")

    assert hammerdown.core.convert_file(source)
    assert (tmp_path / "hm_notes_txt.md").read_text(encoding="utf-8") == "line one\nline two\n"


@pytest.mark.parametrize("removed_flag", ["--in-place", "--in_place", "--skip_existing", "--force", "-f", "-q", "--quiet"])
def test_removed_flags_are_rejected(removed_flag: str):
    with pytest.raises(SystemExit) as exc:
        hammerdown.cli.main([removed_flag])

    assert exc.value.code == 2


def test_short_include_flag_requires_a_pattern():
    with pytest.raises(SystemExit) as exc:
        hammerdown.cli.main(["-i"])

    assert exc.value.code == 2


def test_overwrite_by_default(tmp_path: Path):
    source = tmp_path / "notes.txt"
    source.write_text("line one\n", encoding="utf-8")
    assert hammerdown.core.convert_file(source)
    assert (tmp_path / "hm_notes_txt.md").read_text(encoding="utf-8") == "line one\n"

    # Modifying source and running again overwrites existing output by default
    source.write_text("updated line\n", encoding="utf-8")
    assert hammerdown.cli.main([str(source)]) == 0
    assert (tmp_path / "hm_notes_txt.md").read_text(encoding="utf-8") == "updated line\n"


def test_skip_existing_flag(tmp_path: Path):
    source = tmp_path / "notes.txt"
    source.write_text("line one\n", encoding="utf-8")
    assert hammerdown.cli.main([str(source)]) == 0
    assert (tmp_path / "hm_notes_txt.md").read_text(encoding="utf-8") == "line one\n"

    # Modify source, but with -s / --skip-existing the output is not overwritten
    source.write_text("modified text\n", encoding="utf-8")
    assert hammerdown.cli.main(["-s", str(source)]) == 0
    assert (tmp_path / "hm_notes_txt.md").read_text(encoding="utf-8") == "line one\n"

    assert hammerdown.cli.main(["--skip-existing", str(source)]) == 0
    assert (tmp_path / "hm_notes_txt.md").read_text(encoding="utf-8") == "line one\n"



def test_skip_existing_and_overwrite_statuses(tmp_path: Path, capsys):
    import logging

    source = tmp_path / "notes.txt"
    source.write_text("initial\n", encoding="utf-8")
    assert hammerdown.cli.main([str(source)]) == 0
    source.write_text("updated\n", encoding="utf-8")

    with capsys.disabled():
        assert hammerdown.cli.main(["--skip-existing", str(source)]) == 0
    assert (tmp_path / "hm_notes_txt.md").read_text(encoding="utf-8") == "initial\n"

    with capsys.disabled():
        assert hammerdown.cli.main(["--overwrite", str(source)]) == 0
    assert source.read_text(encoding="utf-8") == "updated\n"


@pytest.mark.parametrize("removed_flag", ["-q", "--quiet"])
def test_removed_quiet_flags_are_rejected(removed_flag: str):
    with pytest.raises(SystemExit) as exc:
        hammerdown.cli.main([removed_flag])

    assert exc.value.code == 2


def test_unsupported_file_returns_failure(tmp_path: Path):
    source = tmp_path / "image.png"
    source.write_bytes(b"image")

    assert hammerdown.cli.main([str(source)]) == 1


def test_noisy_error_diagnostics_are_in_error_table_not_results(tmp_path: Path, capsys):
    missing = tmp_path / "missing.txt"

    assert hammerdown.cli.main([str(missing)]) == 1
    output = capsys.readouterr().out

    assert "Error list (1)" in output
    assert "missing.txt" in output
    assert "Results" in output
    results_section = output.split("Results", 1)[1]
    assert "missing.txt" not in results_section
    assert "Errors: 1" in output


def test_mixed_existing_and_missing_files_continue(tmp_path: Path, capsys):
    import logging

    source = tmp_path / "valid.txt"
    source.write_text("valid content\n", encoding="utf-8")
    missing = tmp_path / "missing.txt"

    assert hammerdown.cli.main([str(missing), str(source)]) == 1

    assert (tmp_path / "hm_valid_txt.md").is_file()
    output = capsys.readouterr().out
    assert "Error list (1)" in output
    assert "missing.txt" in output
    assert "file or directory not found" in output


def test_version_option(capsys):
    try:
        hammerdown.cli.main(["--version"])
    except SystemExit as error:
        assert error.code == 0
    assert f"hammerdown {hammerdown.__version__}" in capsys.readouterr().out


def test_overwrite_flag_plain_text(tmp_path: Path):
    source_txt = tmp_path / "notes.txt"
    source_txt.write_text("initial text\n", encoding="utf-8")

    assert hammerdown.cli.main(["-w", str(source_txt)]) == 0
    assert not (tmp_path / "hm_notes_txt.md").exists()
    assert source_txt.read_text(encoding="utf-8") == "initial text\n"

    source_md = tmp_path / "doc.md"
    source_md.write_text("# Heading\n\nSome text   \n", encoding="utf-8")
    assert hammerdown.cli.main(["--overwrite", str(source_md)]) == 0
    assert not (tmp_path / "hm_doc_md.md").exists()
    assert source_md.read_text(encoding="utf-8") == "# Heading\n\nSome text\n"


def test_overwrite_flag_ignored_for_binary_format(tmp_path: Path):
    import docx

    source_docx = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_paragraph("Sample content")
    doc.save(str(source_docx))

    orig_bytes = source_docx.read_bytes()
    assert hammerdown.cli.main(["-w", str(source_docx)]) == 0

    # Binary source file must remain untouched
    assert source_docx.read_bytes() == orig_bytes
    out_md = tmp_path / "hm_test_docx.md"
    assert out_md.is_file()
    assert "Sample content" in out_md.read_text(encoding="utf-8")


def test_logging_progress_from_zero(tmp_path: Path, caplog):
    import logging
    import docx

    source_docx = tmp_path / "doc.docx"
    doc = docx.Document()
    for i in range(5):
        doc.add_paragraph(f"Paragraph {i}")
    doc.save(str(source_docx))

    with caplog.at_level(logging.INFO, logger="hammerdown"):
        assert hammerdown.core.convert_file(source_docx)

    progress_messages = [record.message for record in caplog.records if "(0%)" in record.message]
    assert len(progress_messages) > 0, "Expected a log message starting at 0%"


def test_recursive_directory_traversal(tmp_path: Path):
    root_dir = tmp_path / "documents"
    sub_dir = root_dir / "nested" / "deep"
    sub_dir.mkdir(parents=True)

    file1 = root_dir / "doc1.txt"
    file1.write_text("content 1\n", encoding="utf-8")

    file2 = sub_dir / "doc2.txt"
    file2.write_text("content 2\n", encoding="utf-8")

    ignored_temp = root_dir / "~$temp.docx"
    ignored_temp.write_bytes(b"temp")

    ignored_hidden = sub_dir / ".hidden.txt"
    ignored_hidden.write_text("secret\n", encoding="utf-8")

    assert hammerdown.cli.main([str(root_dir)]) == 0

    assert (root_dir / "hm_doc1_txt.md").is_file()
    assert (root_dir / "hm_doc1_txt.md").read_text(encoding="utf-8") == "content 1\n"
    assert (sub_dir / "hm_doc2_txt.md").is_file()
    assert (sub_dir / "hm_doc2_txt.md").read_text(encoding="utf-8") == "content 2\n"

    assert not (root_dir / "hm_~$temp_docx.md").exists()
    assert not (sub_dir / "hm_.hidden_txt.md").exists()


def test_directory_converts_supported_files_and_ignores_unsupported(tmp_path: Path):
    root_dir = tmp_path / "flat"
    root_dir.mkdir()
    first = root_dir / "first.txt"
    second = root_dir / "second.md"
    unsupported = root_dir / "notes.bin"
    first.write_text("first document\n", encoding="utf-8")
    second.write_text("# second document\n", encoding="utf-8")
    unsupported.write_bytes(b"not a supported format")

    assert hammerdown.cli.main([str(root_dir)]) == 0

    assert (root_dir / "hm_first_txt.md").read_text(encoding="utf-8") == "first document\n"
    assert (root_dir / "hm_second_md.md").read_text(encoding="utf-8") == "# second document\n"
    assert not (root_dir / "hm_notes_bin.md").exists()


def test_directory_and_explicit_child_file_are_deduplicated(tmp_path: Path, capsys):
    import logging

    root_dir = tmp_path / "deduplicated"
    nested = root_dir / "nested"
    nested.mkdir(parents=True)
    child = nested / "child.txt"
    child.write_text("convert once\n", encoding="utf-8")

    assert hammerdown.cli.main([str(root_dir), str(child)]) == 0

    output = capsys.readouterr().out
    assert (nested / "hm_child_txt.md").read_text(encoding="utf-8") == "convert once\n"
    assert output.count("Completed child.txt") == 0
    assert output.count("converted files") == 1




def test_include_and_exclude_filter_order_and_repeated_patterns(tmp_path: Path, capsys):
    root = tmp_path / "documents"
    nested = root / "nested"
    nested.mkdir(parents=True)
    keep_a = root / "keep_a.txt"
    keep_b = nested / "keep_b.md"
    exclude = nested / "exclude_keep.txt"
    ignore = root / "other.txt"
    for path in (keep_a, keep_b, exclude, ignore):
        path.write_text(path.name, encoding="utf-8")

    assert hammerdown.cli.main([
        "-i", "keep", "-i", "other", "-e", "exclude", str(root)
    ]) == 0
    capsys.readouterr()

    assert (root / "hm_keep_a_txt.md").is_file()
    assert (nested / "hm_keep_b_md.md").is_file()
    assert (nested / "hm_exclude_keep_txt.md").exists() is False
    assert (root / "hm_other_txt.md").is_file()


def test_include_selects_subset_and_quoted_pattern_is_normalized(tmp_path: Path, capsys):
    root = tmp_path / "docs"
    root.mkdir()
    selected = root / "Русское предложение (x).txt"
    other = root / "other.txt"
    selected.write_text("selected", encoding="utf-8")
    other.write_text("other", encoding="utf-8")

    assert hammerdown.cli.main([
        "-i", "'Русское предложение \\(x\\)'", str(root)
    ]) == 0
    capsys.readouterr()

    assert (root / "hm_Русское предложение (x)_txt.md").is_file()
    assert not (root / "hm_other_txt.md").exists()


def test_invalid_filter_regex_is_argparse_error(tmp_path: Path):
    source = tmp_path / "notes.txt"
    source.write_text("text", encoding="utf-8")

    with pytest.raises(SystemExit) as exc:
        hammerdown.cli.main(["--include", "[", str(source)])

    assert exc.value.code == 2
    assert not (tmp_path / "hm_notes_txt.md").exists()


def test_stdout_contains_only_merged_markdown_and_does_not_write_files(tmp_path: Path, capsys):
    first = tmp_path / "one.txt"
    second = tmp_path / "two.txt"
    first.write_text("one content", encoding="utf-8")
    second.write_text("two content", encoding="utf-8")

    assert hammerdown.cli.main(["--stdout", str(first), str(second)]) == 0
    captured = capsys.readouterr()

    assert "FILE: one.txt" in captured.out
    assert "FILE: two.txt" in captured.out
    assert "one content" in captured.out
    assert "two content" in captured.out
    assert captured.err == ""
    assert not (tmp_path / "hm_one_txt.md").exists()
    assert not (tmp_path / "hm_two_txt.md").exists()


def test_stdout_confirm_cancels_without_interactive_input(tmp_path: Path, capsys):
    source = tmp_path / "notes.txt"
    source.write_text("content", encoding="utf-8")

    assert hammerdown.cli.main(["--stdout", "--confirm", str(source)]) == 1
    captured = capsys.readouterr()

    assert captured.out == ""
    assert "Preflight:" in captured.err
    assert "requires interactive stdin" in captured.err
    assert not (tmp_path / "hm_notes_txt.md").exists()


def test_stdout_confirm_only_accepts_y(monkeypatch, tmp_path: Path, capsys):
    import io

    class InteractiveInput(io.StringIO):
        def isatty(self):
            return True

    source = tmp_path / "notes.txt"
    source.write_text("content", encoding="utf-8")
    monkeypatch.setattr("sys.stdin", InteractiveInput("yes\\n"))

    assert hammerdown.cli.main(["--stdout", "--confirm", str(source)]) == 1
    captured = capsys.readouterr()

    assert captured.out == ""
    assert "Cancelled." in captured.err


def test_stdout_confirm_accepts_trimmed_case_insensitive_y(monkeypatch, tmp_path: Path, capsys):
    import io

    class InteractiveInput(io.StringIO):
        def isatty(self):
            return True

    source = tmp_path / "notes.txt"
    source.write_text("content", encoding="utf-8")
    monkeypatch.setattr("sys.stdin", InteractiveInput(" Y \n"))

    assert hammerdown.cli.main(["--stdout", "--confirm", str(source)]) == 0
    captured = capsys.readouterr()

    assert "FILE: notes.txt" in captured.out
    assert "Preflight:" in captured.err
    assert "Continue and write Markdown" in captured.err
    assert not (tmp_path / "hm_notes_txt.md").exists()


def test_enter_wait_is_exclusive_to_right_click(tmp_path: Path, monkeypatch):
    source = tmp_path / "notes.txt"
    source.write_text("content", encoding="utf-8")
    calls = []
    monkeypatch.setattr("builtins.input", lambda: calls.append("wait"))

    assert hammerdown.cli.main([str(source)]) == 0
    assert calls == []
    assert hammerdown.cli.main(["--right-click", str(source)]) == 0
    assert calls == ["wait"]


def test_merge_converted_files(tmp_path: Path):
    file1 = tmp_path / "doc1.txt"
    file1.write_text("Hello from file 1\n", encoding="utf-8")
    file2 = tmp_path / "doc2.txt"
    file2.write_text("Hello from file 2\n", encoding="utf-8")

    merged_file = tmp_path / "summary.md"
    assert hammerdown.cli.main(["--merge", str(merged_file), str(file1), str(file2)]) == 0

    assert merged_file.is_file()
    content = merged_file.read_text(encoding="utf-8")
    assert "FILE: doc1.txt" in content
    assert "FILE: doc2.txt" in content
    assert "# doc1.txt" not in content
    assert "# doc2.txt" not in content
    assert "Hello from file 1" in content
    assert "Hello from file 2" in content


def test_merge_with_short_flag_and_skipped_failure(tmp_path: Path):
    file1 = tmp_path / "doc1.txt"
    file1.write_text("Hello from file 1\n", encoding="utf-8")
    bad_file = tmp_path / "missing.docx"

    merged_file = tmp_path / "partial.md"
    assert hammerdown.cli.main(["-m", str(merged_file), str(file1), str(bad_file)]) == 1

    assert merged_file.is_file()
    content = merged_file.read_text(encoding="utf-8")
    assert "FILE: doc1.txt" in content
    assert "missing.docx" not in content
